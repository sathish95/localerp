"""
styles.py — the shared Arial + blue/black/green/yellow convention every
tab in both build_costing.py and build_roi.py uses. Not part of the
original spec's file list, but factored out so the convention can't drift
between the two build scripts (a genuine two-line change per file
otherwise, repeated across a dozen tabs).

Convention:
  blue   = manual input (a constant from inputs.py, entered directly on
           the Inputs tab)
  black  = formula (computed from other cells on the SAME sheet)
  green  = cross-sheet link (a formula that just references another
           sheet's cell, e.g. ='Inputs'!B4)
  yellow fill + comment = a value inputs.py could not supply — never a
           silently invented number; the workbook makes that visible at
           the cell itself, not just in a code comment.
"""
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment

ARIAL = 'Arial'
BLUE = 'FF0000FF'
BLACK = 'FF000000'
GREEN = 'FF008000'
BOLD_BLACK = 'FF000000'

YELLOW_FILL = PatternFill(start_color='FFFFF200', end_color='FFFFF200', fill_type='solid')
HEADER_FILL = PatternFill(start_color='FFD9D9D9', end_color='FFD9D9D9', fill_type='solid')

CURRENCY_FMT = '"₹"#,##0.00'
PERCENT_FMT = '0.0%'
NUMBER_FMT = '#,##0'

NEEDS_INPUT_NOTE = 'NEEDS INPUT FROM SATHISHKUMAR'


def _font(color, bold=False, size=10):
    return Font(name=ARIAL, color=color, bold=bold, size=size)


def as_input(cell, number_format=CURRENCY_FMT):
    """A manual constant from inputs.py, entered on the Inputs tab."""
    cell.font = _font(BLUE)
    cell.number_format = number_format
    return cell


def as_formula(cell, number_format=CURRENCY_FMT, bold=False):
    """A same-sheet computed formula."""
    cell.font = _font(BLACK, bold=bold)
    cell.number_format = number_format
    return cell


def as_link(cell, number_format=CURRENCY_FMT):
    """A cross-sheet reference formula (e.g. ='Inputs'!B4)."""
    cell.font = _font(GREEN)
    cell.number_format = number_format
    return cell


def flag_needs_input(cell, note=NEEDS_INPUT_NOTE):
    """Marks a placeholder value that inputs.py could not supply for real."""
    cell.fill = YELLOW_FILL
    cell.comment = Comment(note, 'Costing Generator')
    return cell


def header(cell, text):
    cell.value = text
    cell.font = _font(BLACK, bold=True, size=11)
    cell.fill = HEADER_FILL
    return cell
