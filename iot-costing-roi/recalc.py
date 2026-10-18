"""
recalc.py — forces LibreOffice Calc to recompute every formula in the
generated workbook, then scans the recalculated file for formula errors
and fails loudly (non-zero exit, printed cell list) if any are found.

Why via LibreOffice and not just trust openpyxl's cached values: openpyxl
writes formula *strings*, not computed results — Excel/LibreOffice compute
the actual numbers the first time a human (or a recalculation pass) opens
the file. tests/test_totals.py needs real computed numbers to assert
against, so this step has to run between build.py writing the file and
the tests loading it.

Requires a local LibreOffice install (the `soffice` binary) — this
environment did not have one installed when this repo was scaffolded, so
this script's `soffice --headless --convert-to` invocation is written to
the standard, widely-documented pattern but has NOT been run end-to-end
here. Run `python recalc.py <path-to-workbook.xlsx>` yourself once
LibreOffice is installed and tell me if anything about the invocation
needs adjusting for your OS/LibreOffice version.
"""
import shutil
import subprocess
import sys
import tempfile
import os
from pathlib import Path

from openpyxl import load_workbook

ERROR_STRINGS = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NULL!', '#NUM!'}

# Common install locations Windows doesn't always put on PATH.
WINDOWS_CANDIDATES = [
    r'C:\Program Files\LibreOffice\program\soffice.exe',
    r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
]


def find_soffice():
    found = shutil.which('soffice') or shutil.which('soffice.exe')
    if found:
        return found
    for candidate in WINDOWS_CANDIDATES:
        if os.path.exists(candidate):
            return candidate
    return None


def recalculate(xlsx_path: str) -> str:
    """Recalculates xlsx_path in place via a LibreOffice headless round-trip. Returns the path (unchanged)."""
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            'LibreOffice (soffice) was not found on PATH or in the usual Windows install '
            'locations. Install LibreOffice (https://www.libreoffice.org/download/) and '
            're-run, or run build.py with --skip-recalc if you only need the formulas '
            'written (not verified) for now.'
        )

    xlsx_path = str(Path(xlsx_path).resolve())
    with tempfile.TemporaryDirectory() as tmpdir:
        # A fresh, isolated LibreOffice user profile avoids picking up any
        # machine-specific settings (e.g. a "never recalculate" override)
        # that would silently defeat the whole point of this script.
        profile_dir = Path(tmpdir) / 'lo_profile'
        cmd = [
            soffice, '--headless', '--norestore',
            f'-env:UserInstallation=file:///{profile_dir.as_posix()}',
            '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
            '--outdir', tmpdir, xlsx_path,
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        if result.returncode != 0:
            raise RuntimeError(f'LibreOffice headless conversion failed:\n{result.stdout}\n{result.stderr}')

        converted = Path(tmpdir) / Path(xlsx_path).name
        if not converted.exists():
            raise RuntimeError(f'Expected recalculated file not found at {converted} — LibreOffice output:\n{result.stdout}\n{result.stderr}')
        shutil.copyfile(converted, xlsx_path)

    return xlsx_path


def check_for_errors(xlsx_path: str):
    """Scans every cell of every sheet for a cached Excel error value. Raises with a full list if any are found."""
    wb = load_workbook(xlsx_path, data_only=True)
    errors = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERROR_STRINGS:
                    errors.append(f'{sheet.title}!{cell.coordinate} = {cell.value}')
    if errors:
        raise RuntimeError('Formula errors found after recalculation:\n' + '\n'.join(errors))


def main():
    if len(sys.argv) != 2:
        print('Usage: python recalc.py <path-to-workbook.xlsx>')
        sys.exit(1)
    path = sys.argv[1]
    recalculate(path)
    check_for_errors(path)
    print(f'Recalculated and verified with no formula errors: {path}')


if __name__ == '__main__':
    main()
